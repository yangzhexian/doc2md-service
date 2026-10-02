"""Regenerate the tiny, self-authored corpus; no external documents or fonts.

Run using the locked project environment: python tests/fixtures/real_documents/generate.py
Only fixture creation needs python-docx, python-pptx, openpyxl, and Pillow.
Regression tests consume the committed files and never run this generator.
"""

from __future__ import annotations

import hashlib
import importlib.metadata
import json
import re
import struct
import zipfile
import zlib
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Inches as DocxInches
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent
FIXED_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)
HTML_NAME = "article-about-concurrency-and-correctness-with-local-resources.html"


def normalized_zip(path: Path) -> None:
    """Remove machine/time-dependent ZIP metadata and sort entries."""
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    if "docProps/core.xml" in entries:
        # openpyxl sets modified to the current time while saving.
        entries["docProps/core.xml"] = re.sub(
            rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
            rb"\g<1>2026-01-01T00:00:00Z\g<2>",
            entries["docProps/core.xml"],
        )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in sorted(entries.items()):
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            archive.writestr(info, data)


def png_from_rgb(width: int, height: int, pixels: bytes) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(
            ">I", zlib.crc32(kind + data) & 0xFFFFFFFF
        )

    rows = b"".join(b"\x00" + pixels[row * width * 3 : (row + 1) * width * 3]
                    for row in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(rows, level=9)) + chunk(b"IEND", b""))


def create_chart() -> tuple[int, int, bytes]:
    width, height = 160, 100
    pixels = bytearray()
    for y in range(height):
        for x in range(width):
            if 15 <= x < 60 and 35 <= y < 85:
                color = (32, 96, 192)
            elif 90 <= x < 135 and 15 <= y < 85:
                color = (224, 96, 32)
            else:
                color = (255, 255, 255)
            pixels.extend(color)
    (ROOT / "assets" / "local-chart.png").write_bytes(png_from_rgb(width, height, pixels))
    return width, height, bytes(pixels)


def create_pdf(width: int, height: int, pixels: bytes) -> None:
    """Write a real PDF with native Helvetica text, two pages and an image."""
    def stream(header: bytes, data: bytes) -> bytes:
        return header + f" /Length {len(data)} >>\nstream\n".encode() + data + b"\nendstream"

    def page_content(title: str, token: str, image: bool) -> bytes:
        lines = [title, token, "This self-authored PDF contains selectable native text.",
                 "Regression checks preserve words, page ranges, and local images.",
                 "Each page has a unique marker so excluded pages cannot pass.",
                 "No web access, scanned text, or downloaded model is required."]
        commands = [b"BT /F1 14 Tf 50 750 Td 24 TL"]
        for index, line in enumerate(lines):
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(("T* " if index else "").encode() + f"({escaped}) Tj".encode())
        commands.append(b"ET")
        if image:
            commands.append(b"q 240 0 0 150 50 400 cm /Im1 Do Q")
        return b"\n".join(commands)

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> /XObject << /Im1 8 0 R >> >> /Contents 6 0 R >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 7 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        stream(b"<<", page_content("Docs2md PDF first page", "FIRSTPAGEALPHA", True)),
        stream(b"<<", page_content("Docs2md PDF second page", "SECONDPAGEOMEGA", False)),
        stream(f"<< /Type /XObject /Subtype /Image /Width {width} /Height {height} /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode".encode(), zlib.compress(pixels, 9)),
    ]
    pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    (ROOT / "two-pages.pdf").write_bytes(pdf)


def create_office() -> None:
    doc = Document()
    doc.core_properties.created = FIXED_TIME
    doc.core_properties.modified = FIXED_TIME
    doc.core_properties.author = "docs2md fixture generator"
    doc.add_heading("Docs2md DOCX sample", level=1)
    doc.add_paragraph("ParagraphTokenDOCX: durable conversions preserve 中文回归.")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("Item", "Quantity", "Note")):
        cell.text = value
    for values in (("Apricot", "17", "中文回归"), ("Walnut", "23", "TableTokenDOCX")):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    doc.add_picture(str(ROOT / "assets" / "local-chart.png"), width=DocxInches(2))
    doc.save(ROOT / "paragraph-and-table.docx")
    normalized_zip(ROOT / "paragraph-and-table.docx")

    deck = Presentation()
    deck.core_properties.created = FIXED_TIME
    deck.core_properties.modified = FIXED_TIME
    deck.core_properties.author = "docs2md fixture generator"
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Docs2md slide heading"
    text_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(1))
    text_box.text = "SlideTokenPPTX: reliable document conversion"
    slide.shapes.add_picture(str(ROOT / "assets" / "local-chart.png"), Inches(0.5), Inches(3), width=Inches(2))
    deck.save(ROOT / "heading-and-image.pptx")
    normalized_zip(ROOT / "heading-and-image.pptx")

    workbook = Workbook()
    workbook.properties.created = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.creator = "docs2md fixture generator"
    worksheet = workbook.active
    worksheet.title = "Inventory"
    for row in (("Item", "Quantity", "Note"), ("Apricot", 17, "中文回归"), ("Walnut", 23, "SheetTokenXLSX")):
        worksheet.append(row)
    workbook.save(ROOT / "inventory.xlsx")
    normalized_zip(ROOT / "inventory.xlsx")


def create_text_and_scan() -> None:
    (ROOT / "inventory.csv").write_bytes(
        'Item,Quantity,Note\nApricot,17,中文回归\nWalnut,23,"Comma, inside quoted cell"\n'.encode("utf-8")
    )
    (ROOT / HTML_NAME).write_bytes("""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Docs2md HTML fixture</title></head>
<body><h1>Local resources survive long filenames</h1>
<p>HTMLTokenResources: preserve this content and the linked image.</p>
<p><a href="https://example.invalid/docs2md-regression">External documentation</a></p>
<img src="assets/local-chart.png" alt="Self-authored two-bar chart">
<table><tr><th>Item</th><th>Quantity</th></tr><tr><td>Apricot</td><td>17</td></tr></table>
</body></html>
""".encode("utf-8"))
    image = Image.new("RGB", (1400, 480), "white")
    draw = ImageDraw.Draw(image)
    # Pillow ships this font. It avoids operating-system font lookup.
    font = ImageFont.load_default(size=56)
    draw.text((70, 90), "DOCS2MD SCANNED SAMPLE", font=font, fill="black")
    draw.text((70, 205), "OCR TOKEN ALPHA 7301", font=font, fill="black")
    image.save(ROOT / "scanned-text.png", compress_level=9)


def create_manifest() -> None:
    cases = [
        {"id": "pdf_all", "file": "two-pages.pdf", "engines": ["markitdown", "mineru"], "contains": ["FIRSTPAGEALPHA", "SECONDPAGEOMEGA"], "mineru_min_images": 1},
        {"id": "pdf_first", "file": "two-pages.pdf", "engines": ["mineru"], "pages": "1", "contains": ["FIRSTPAGEALPHA"], "excludes": ["SECONDPAGEOMEGA"], "mineru_min_images": 1},
        {"id": "pdf_last", "file": "two-pages.pdf", "engines": ["mineru"], "pages": "r1", "contains": ["SECONDPAGEOMEGA"], "excludes": ["FIRSTPAGEALPHA"]},
        {"id": "docx", "file": "paragraph-and-table.docx", "engines": ["markitdown", "mineru"], "contains": ["Docs2md DOCX sample", "ParagraphTokenDOCX", "Apricot", "17", "Walnut", "23", "TableTokenDOCX", "中文回归"], "table_cells": ["Apricot", "17", "Walnut", "23"], "mineru_min_images": 1},
        {"id": "pptx", "file": "heading-and-image.pptx", "engines": ["markitdown", "mineru"], "contains": ["Docs2md slide heading", "SlideTokenPPTX"], "mineru_min_images": 1},
        {"id": "xlsx", "file": "inventory.xlsx", "engines": ["markitdown", "mineru"], "contains": ["Apricot", "17", "Walnut", "23", "SheetTokenXLSX", "中文回归"], "table_cells": ["Apricot", "17", "Walnut", "23"]},
        {"id": "html_long_local", "file": HTML_NAME, "engines": ["markitdown", "mineru"], "contains": ["HTMLTokenResources", "Apricot", "17"], "links": ["https://example.invalid/docs2md-regression"], "table_cells": ["Apricot", "17"], "mineru_min_images": 1, "markitdown_image_target": "assets/local-chart.png"},
        {"id": "csv", "file": "inventory.csv", "engines": ["markitdown", "mineru"], "contains": ["Apricot", "17", "Walnut", "23", "Comma, inside quoted cell", "中文回归"], "mineru_table_cells": ["Apricot", "17", "Walnut", "23"]},
        {"id": "png_ocr", "file": "scanned-text.png", "engines": ["mineru"], "requires_models": True, "contains": ["DOCS2MD", "SCANNED", "OCR", "ALPHA", "7301"]},
    ]
    files = sorted({case["file"] for case in cases} | {"assets/local-chart.png"})
    manifest = {
        "schema_version": 1,
        "provenance": "Self-authored synthetic documents generated by generate.py; no third-party source material, network downloads, or OS fonts.",
        "license": "MIT (same as the repository)",
        "generator_versions": {name: importlib.metadata.version(name) for name in ("python-docx", "python-pptx", "openpyxl", "pillow")},
        "files": {name: {"sha256": hashlib.sha256((ROOT / name).read_bytes()).hexdigest(), "bytes": (ROOT / name).stat().st_size} for name in files},
        "cases": cases,
    }
    (ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")


def main() -> None:
    (ROOT / "assets").mkdir(exist_ok=True)
    create_pdf(*create_chart())
    create_office()
    create_text_and_scan()
    create_manifest()


if __name__ == "__main__":
    main()
